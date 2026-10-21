# -*- coding: utf-8 -*-
"""Supplementary Table 3 — the differential expression lists and gene sets used as references (Methods, "Differential expression
lists and gene sets used as references").

  <python> figs/Supp_Table3_References/analysis/supp_table_references.py        (from the project root; pandas, openpyxl)

Input : figs/Supp_Table3_References/output/SuppTable3_reference_sizes.csv (supp_table_references.R: the number of genes of each
        reference before and after the restriction to the genes analysed, computed from the reference files used for scoring)
Output: figs/Supp_Table3_References/output/SuppTable3_references.xlsx (README + sheet "References")
The descriptions below follow the scripts that build the references: analysis/hammond/06_build_references.R (DAM, Maturity,
Interferon, G2M), analysis/dedonno/03_references.R (Glucocorticoid, Seizure, NF-kB, UPR),
figs/Fig5_benchmark/analysis/fig5_extra_refs.R (PS2APP, Forced swim) and analysis/kang/03_references.R (IFNb, PGE2, Interferon
with human genes, Chemokine).
"""
import datetime, os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = "figs/Supp_Table3_References/output"
F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HEAD = PatternFill("solid", fgColor="DDE3EA")

n = pd.read_csv(os.path.join(OUT, "SuppTable3_reference_sizes.csv")).set_index("reference")

DATASET = {"Hammond": "Hammond et al. (Fig. 2)", "DeDonno": "De Donno et al., wild-type mice (Fig. 3)",
           "Kang": "Kang et al., human PBMC (Supplementary Fig. S5)"}
TYPE = {"B": "A (differential expression list)", "C": "B (gene set)"}
EXCL_A = "Genes without informative names, mitochondrial genes and haemoglobin genes"
BULK = "Counts summed per group, converted to CPM over the nuclear genes; genes ranked by the difference in log2(CPM + 1)"
HALLMARK = "MSigDB Hallmark {}, mouse orthologues (msigdbr 26.1.0; MSigDB 2026.1.Hs)"
SET_RULE = "Membership: the whole set, without order or direction"
KANG_NAMES = "; gene names matched to the data of Kang et al. by Ensembl ID"
## one row per reference, in the order of Methods: the six differential expression lists, then the four gene sets; then the four
## references of Supplementary Fig. S5 (human)
ROWS = [
    ("DAM", "Figs. 2c, e, f and 5b, c, e",
     "Keren-Shaul et al., 2017; Table S3 of the original article",
     "Disease-associated microglia (DAM) vs homeostatic microglia (single-cell RNA-seq of 5XFAD mice)", "Published fold changes",
     "Published fold change (genes with a finite, non-zero fold change); top 500 per direction, then restricted to the genes analysed", EXCL_A),
    ("Maturity", "Fig. 2c, e; Supplementary Table 2",
     "Matcovitch-Natan et al., 2016; GEO GSE79812 (supplementary table GSE79812_Timeline.xlsx)",
     "Adult microglia (cortex and hippocampus) vs embryonic microglia (E13.5, E14.5 and E16.5)", "4 vs 4 samples",
     "Normalized values of the GEO table summed per group and scaled to one million over the nuclear genes; genes ranked by the "
     "difference in log2(value + 1); top 500 per direction, then restricted to the genes analysed", EXCL_A),
    ("PS2APP", "Fig. 5e; Supplementary Fig. S4",
     "Srinivasan et al., 2016; EMBL-EBI Expression Atlas E-GEOD-75431 (= GEO GSE75431), raw counts per gene",
     "Sorted cortical microglia of PS2APP vs non-transgenic mice, 13 months", "5 vs 5 mice",
     BULK + "; top 500 per direction, then restricted to the genes analysed", EXCL_A),
    ("Glucocorticoid", "Figs. 3c, e and 5c, e",
     "Viho et al., 2022; GEO GSE184924 (gene-level counts)",
     "Ventral hippocampus 3 h after corticosterone (3 mg/kg) vs vehicle injection", "4 vs 4 mice",
     BULK + "; top 500 per direction, then restricted to the genes analysed", EXCL_A),
    ("Forced swim", "Fig. 5e; Supplementary Fig. S4",
     "von Ziegler et al., 2022; GEO GSE169505 (gene-level counts and series matrices)",
     "Ventral hippocampus of male mice 3 h after a 6-min forced swim vs home cage (experiment TimeSeriesFST)", "7 vs 8 mice",
     BULK + "; top 500 per direction, then restricted to the genes analysed", EXCL_A),
    ("Seizure", "Fig. 3c, e",
     "Fernandez-Albert et al., 2019; GEO GSE125068 via recount3 (SRP179090; gene-level base-pair coverage counts)",
     "Nuclear RNA of sorted nuclei of hippocampal excitatory neurons 6 h after kainate-induced seizures (status epilepticus; "
     "kainic acid 25 mg/kg) vs saline (GSM3562104-05 vs GSM3562100-01)", "2 vs 2 samples",
     BULK + "; up-regulated genes only: the top 200, then restricted to the genes analysed", EXCL_A),
    ("Interferon", "Figs. 2d, f and 5d", HALLMARK.format("INTERFERON_ALPHA_RESPONSE"), "—", "—", SET_RULE, "—"),
    ("G2M", "Figs. 2d and 5d", HALLMARK.format("G2M_CHECKPOINT"), "—", "—", SET_RULE, "—"),
    ("NF-κB", "Fig. 3d, f", HALLMARK.format("TNFA_SIGNALING_VIA_NFKB"), "—", "—", SET_RULE, "—"),
    ("UPR", "Fig. 3d, f", HALLMARK.format("UNFOLDED_PROTEIN_RESPONSE"), "—", "—", SET_RULE, "—"),
    ("IFNb", "Supplementary Fig. S5c, e",
     "Galbraith et al., 2022; GEO GSE191317 (raw gene counts; 3′ mRNA-seq)",
     "Peripheral blood mononuclear cells, IFN-β vs mock (PBS), 18 h", "7 vs 7 donors",
     BULK + "; top 500 per direction, then restricted to the genes analysed" + KANG_NAMES, EXCL_A),
    ("PGE2", "Supplementary Fig. S5c, e",
     "Sokhi et al., 2025; GEO GSE272019 (raw gene counts)",
     "CD14+ monocytes cultured overnight with M-CSF, PGE2 (280 nM) vs unstimulated, 3 h", "3 vs 3 donors",
     BULK + "; top 500 per direction, then restricted to the genes analysed" + KANG_NAMES, EXCL_A),
    ("Interferon (human)", "Supplementary Fig. S5d, f",
     "MSigDB Hallmark INTERFERON_ALPHA_RESPONSE, human genes (msigdbr 26.1.0; MSigDB 2026.1.Hs)" + KANG_NAMES, "—", "—", SET_RULE, "—"),
    ("Chemokine", "Supplementary Fig. S5d, f",
     "KEGG CHEMOKINE_SIGNALING_PATHWAY (MSigDB C2, KEGG legacy), human genes (msigdbr 26.1.0; MSigDB 2026.1.Hs)" + KANG_NAMES,
     "—", "—", SET_RULE, "—"),
]
KEY = {"NF-κB": "NF-kB"}                                       # name in the CSV (ASCII)
assert sorted(KEY.get(r[0], r[0]) for r in ROWS) == sorted(n.index), "references in the CSV and in this script differ"

recs = []
for name, used, src, comp, samples, rank, excl in ROWS:
    c = n.loc[KEY.get(name, name)]
    gs = c.type_code == "C" or c.down_before == 0                 # no down-regulated list: gene sets and Seizure
    recs.append({"Reference": name, "Type": TYPE[c.type_code], "Used in": used, "Cells scored": DATASET[c.dataset],
                 "Source": src, "Comparison (up = higher in the first group)": comp, "Samples": samples,
                 "Ranking or selection of genes": rank, "Genes excluded": excl,
                 "Genes in the reference: up": int(c.up_before), "Genes in the reference: down": None if gs else int(c.down_before),
                 "Genes analysed in the dataset": int(c.genes_analysed),
                 "Genes used, after restriction to the genes analysed: up": int(c.up_after),
                 "Genes used, after restriction to the genes analysed: down": None if gs else int(c.down_after)})
t = pd.DataFrame(recs)

wb = Workbook(); ws = wb.active; ws.title = "README"
readme = [
    ("Supplementary Table 3 | Differential expression lists and gene sets used as references.", FB),
    ("", F),
    ("One row per reference (Methods). Type A references are differential expression lists, ranked by the change of each gene between "
     "two groups and divided into up-regulated (higher in the first group of the comparison) and down-regulated genes; type B "
     "references are gene sets, used whole without order or direction. Each reference was restricted to the genes analysed in the "
     "dataset whose cells it scores before matching; the numbers of genes after this restriction are those used.", F),
    ("", F),
    ("Columns", FB),
    ("Type — A, differential expression list; B, gene set (as in Fig. 5).", F),
    ("Used in — the figures and tables that show the index computed with the reference.", F),
    ("Cells scored — the dataset whose cells were scored with the reference (De Donno et al., the wild-type mice only).", F),
    ("Source — the study and the data from which the reference was taken or computed.", F),
    ("Comparison (up = higher in the first group) — the two groups compared; up-regulated genes are those higher in the first group.", F),
    ("Samples — the numbers of mice, donors or samples in the first and second groups, whose counts were summed per group.", F),
    ("Ranking or selection of genes — how the genes were ranked and how many were taken. CPM, counts per million.", F),
    ("Genes excluded — genes removed before ranking. Genes without informative names: predicted genes, RIKEN cDNA clones and other "
     "clone names (patterns in the analysis code). Gene sets: no genes were removed beyond the restriction to the genes analysed.", F),
    ("Genes in the reference: up / down — number of genes in the up- and down-regulated lists (gene sets: the number of genes in "
     "the set; mouse orthologues for Figs. 2 and 3, human genes for Supplementary Fig. S5) before the restriction to the genes analysed.", F),
    ("Genes analysed in the dataset — the genes analysed in the dataset whose cells were scored (Methods).", F),
    ("Genes used, after restriction to the genes analysed: up / down — number of genes of the reference used for matching. "
     "Empty cells, no down-regulated list (Seizure and the gene sets; the Seizure-index is Score(UP, UP) − Score(DOWN, UP)).", F),
    ("", F),
    (f"Source: figs/Supp_Table3_References/output/SuppTable3_reference_sizes.csv (figs/Supp_Table3_References/analysis/"
     f"supp_table_references.R, which restricts the reference files used for scoring as the scoring scripts do); table made by "
     f"figs/Supp_Table3_References/analysis/supp_table_references.py on {datetime.date.today().isoformat()}.", F),
]
for i, (text, font) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=text); c.font = font; c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140

s = wb.create_sheet("References")
for j, name in enumerate(t.columns, start=1):
    c = s.cell(row=1, column=j, value=name); c.font = FB; c.fill = HEAD
    c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
for i, row in enumerate(t.itertuples(index=False), start=2):
    for j, v in enumerate(row, start=1):
        v = None if (v is None or (isinstance(v, float) and pd.isna(v))) else (v.item() if hasattr(v, "item") else v)
        if isinstance(v, float) and v.is_integer(): v = int(v)
        c = s.cell(row=i, column=j, value=v); c.font = F
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if isinstance(v, int): c.number_format = "#,##0"
for j, w in enumerate([14, 18, 18, 20, 34, 38, 12, 44, 30, 11, 11, 12, 14, 14], start=1):
    s.column_dimensions[get_column_letter(j)].width = w
s.row_dimensions[1].height = 60; s.freeze_panes = "B2"
path = os.path.join(OUT, "SuppTable3_references.xlsx"); wb.save(path)
print("wrote", path)
print(t[["Reference", "Type", "Genes in the reference: up", "Genes in the reference: down", "Genes analysed in the dataset",
         "Genes used, after restriction to the genes analysed: up",
         "Genes used, after restriction to the genes analysed: down"]].to_string(index=False))
