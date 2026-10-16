# -*- coding: utf-8 -*-
"""Supplementary Table 2 — the placement of the five genes of Fig. 4d compared with random sets drawn from all knock-outs of the same cell type.

  <python> figs/Fig4_Perturbseq/analysis/supp_table_random_sets.py        (from the project root; pandas, openpyxl)

Input : figs/Fig4_Perturbseq/output/geneset_random_null.csv (fig4_random_set_null.py),
        figs/Fig4_Perturbseq/output/fig4_other_methods_summary.csv (fig4_other_methods.py; sheet "Other matching methods")
Output: figs/Fig4_Perturbseq/output/SuppTable_random_sets.xlsx (README + two sheets)
"""
import datetime, os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = "figs/Fig4_Perturbseq/output"
F, FB = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
HEAD = PatternFill("solid", fgColor="DDE3EA")

d = pd.read_csv(os.path.join(OUT, "geneset_random_null.csv"))
t = pd.DataFrame({"Cell type": d.cell_type, "Knock-outs (n)": d.n_knockouts, "AUC of the five genes": d.auc_fig4d,
                  "Random sets: mean AUC": d.null_any_mean, "Random sets: SD of AUC": d.null_any_sd, "P, random sets": d.p_any,
                  "Matched sets: mean AUC": d.null_matched_mean, "Matched sets: SD of AUC": d.null_matched_sd,
                  "P, matched sets": d.p_matched})
## sheet "Other matching methods": the same comparison with the knock-outs ranked by ULM and the scFOCAL-type score
## (fig4_other_methods.py; its ULM_all is not in the table); positions = % from the opposed end, as in Fig. 4e
GENES = ["Nr3c1", "Nr3c2", "Hsp90ab1", "Ncoa1", "Ncoa3"]
METH = {"RFind": "RFind-sc", "ULM": "ULM", "scFOCAL": "scFOCAL-type"}
o = pd.read_csv(os.path.join(OUT, "fig4_other_methods_summary.csv"))
o = o[o.method.isin(list(METH))]
o = o.assign(k=o.cell_type.map({c: i for i, c in enumerate(d.cell_type)}) * 10 + o.method.map({m: i for i, m in enumerate(METH)}))
o = o.sort_values("k")
u = pd.DataFrame({"Cell type": o.cell_type, "Matching method": o.method.map(METH), "Knock-outs (n)": o.n_ko,
                  **{f"{g}: % from the opposed end": 100 * o[f"pos_{g}"] / o.n_ko for g in GENES},
                  "AUC of the five genes": o.auc_fig4d, "P, random sets": o.p_any, "P, matched sets": o.p_matched})
rf = o[o.method == "RFind"].set_index("cell_type")   # the rows of RFind-sc must repeat sheet "Random sets" (CSVs of 4-6 digits)
for r in d.itertuples():
    assert rf.loc[r.cell_type, "n_ko"] == r.n_knockouts and abs(rf.loc[r.cell_type, "auc_fig4d"] - r.auc_fig4d) < 1e-4 and \
        abs(rf.loc[r.cell_type, "p_any"] / r.p_any - 1) < 1e-5 and abs(rf.loc[r.cell_type, "p_matched"] / r.p_matched - 1) < 1e-5, r.cell_type
FMT = {"AUC of the five genes": "0.000", "Random sets: mean AUC": "0.000", "Random sets: SD of AUC": "0.000",
       "Matched sets: mean AUC": "0.000", "Matched sets: SD of AUC": "0.000", "P, random sets": "0.00E+00", "P, matched sets": "0.00E+00",
       **{f"{g}: % from the opposed end": "0.0" for g in GENES}}

wb = Workbook(); ws = wb.active; ws.title = "README"
readme = [
    ("Supplementary Table 2 | Comparison with sets of knock-outs drawn at random from all knock-outs of the same cell type.", FB),
    ("", F),
    ("Because the genes targeted by the Perturb-seq library had been selected for broad neuronal expression and relevance to disease (Shi et al., 2026), the "
     "positions of the five knock-outs of Fig. 4d (Nr3c1, Nr3c2, Hsp90ab1, Ncoa1, Ncoa3) were compared with those of sets of five knock-outs drawn at random from all "
     "knock-outs of the same cell type. In each cell type, the AUC of the RFind scores of the five knock-outs against those of the remaining knock-outs "
     "(the rankings of Supplementary Table 1) was compared with the AUCs of 100,000 random sets of five knock-outs, and of 100,000 sets matched to the five genes for the "
     "quartiles of target expression (CPM in the non-targeting nuclei) and of the number of nuclei.", F),
    ("", F),
    ("Sheet \"Other matching methods\" — the same comparison with the knock-outs ranked against the stress response by two other matching methods instead of "
     "RFind-sc (Methods): ULM of decoupleR (the t-value of the slope of the linear regression of the Delta of the stress response on weights equal to the Delta of "
     "the knock-out for its top and bottom 500 genes and 0 for the other genes) and the scFOCAL-type score (the Spearman correlation between the Delta of the "
     "knock-out and that of the stress response over the top and bottom 500 genes of the knock-out). Every method has the same random and matched sets; the rows "
     "of RFind-sc repeat sheet \"Random sets\".", F),
    ("", F),
    ("Columns", FB),
    ("Knock-outs (n) — knock-outs ranked in the cell type.", F),
    ("Matching method — the score by which the knock-outs were ranked (sheet \"Other matching methods\").", F),
    ("<gene>: % from the opposed end — position of the knock-out counted from the knock-out most opposed to the stress response (1), as a percentage of the knock-outs ranked (as in Fig. 4e; below 50, the opposed side).", F),
    ("AUC of the five genes — probability that one of the five knock-outs has a higher RFind score (in sheet \"Other matching methods\", the score of the matching method) than a remaining knock-out (0.5, no bias; near 0, the opposed end; near 1, the same-direction end).", F),
    ("Random sets / Matched sets: mean and SD of AUC — the AUCs of the 100,000 random or matched sets of five knock-outs.", F),
    ("P — fraction of the random (or matched) sets at least as far towards the opposed end as the five genes: (1 + number of sets with AUC <= that of the five genes) / (1 + 100,000).", F),
    ("", F),
    (f"Source: figs/Fig4_Perturbseq/output/geneset_random_null.csv (figs/Fig4_Perturbseq/analysis/fig4_random_set_null.py; seed 1) and "
     f"fig4_other_methods_summary.csv (figs/Fig4_Perturbseq/analysis/fig4_other_methods.py; the same draws); table made by "
     f"figs/Fig4_Perturbseq/analysis/supp_table_random_sets.py on {datetime.date.today().isoformat()}.", F),
]
for i, (text, font) in enumerate(readme, start=1):
    c = ws.cell(row=i, column=1, value=text); c.font = font; c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 140
def sheet(title, t, widths, height=30):
    s = wb.create_sheet(title)
    for j, name in enumerate(t.columns, start=1):
        c = s.cell(row=1, column=j, value=name); c.font = FB; c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for i, row in enumerate(t.itertuples(index=False), start=2):
        for j, (name, v) in enumerate(zip(t.columns, row), start=1):
            c = s.cell(row=i, column=j, value=(v.item() if hasattr(v, "item") else v)); c.font = F
            if name in FMT: c.number_format = FMT[name]
    for j, w in enumerate(widths, start=1):
        s.column_dimensions[get_column_letter(j)].width = w
    s.row_dimensions[1].height = height; s.freeze_panes = "A2"
sheet("Random sets", t, [10, 12, 12, 13, 13, 12, 13, 13, 12])
sheet("Other matching methods", u, [10, 13, 12, 12, 12, 12, 12, 12, 12, 12, 12], height=45)
path = os.path.join(OUT, "SuppTable_random_sets.xlsx"); wb.save(path)
print("wrote", path); print(t.to_string(index=False)); print(u.round(4).to_string(index=False))
